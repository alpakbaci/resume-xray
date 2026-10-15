"""Generate the synthetic test resumes in tests/fixtures/synthetic/.

Every person, company and contact detail here is invented (example.com addresses, 555 numbers).
The files are made with real tools so that they carry real-world quirks:

- PDFs from HTML are printed by Google Chrome (headless), the engine behind many web resume
  builders' "Download PDF" buttons.
- DOCX files are built with python-docx; their PDFs are exported by Microsoft Word.
- "Scanned" PDFs are page images rendered with pypdfium2 and saved by Pillow.

Requires macOS with Google Chrome and Microsoft Word installed. Run from the repository root:

    python tools/make_fixtures.py

The generated files are committed, so the tests do not need these tools.
"""

import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import docx
import pypdfium2
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "tests" / "fixtures" / "synthetic"
FONTS = ROOT / "resume_xray" / "web" / "static" / "fonts"
CHROME = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"
# Word for Mac is sandboxed; it can open and save files in its own container without prompts.
WORD_DIR = Path.home() / "Library/Containers/com.microsoft.Word/Data/Documents/resume-xray-fixtures"

EN = {
    "name": "Alex Morgan",
    "title": "Product Manager",
    "email": "alex.morgan@example.com",
    "phone": "+1 555 010 2030",
    "linkedin": "linkedin.com/in/alex-morgan-example",
    "headings": {"summary": "Summary", "experience": "Experience", "education": "Education", "skills": "Skills", "languages": "Languages"},
    "summary": "Product manager with eight years of experience shipping web and mobile products, from discovery to launch.",
    "jobs": [
        ("Senior Product Manager", "Northwind Apps", "Jan 2021 – Present", [
            "Owned the roadmap for a mobile app with two million monthly users.",
            "Ran A/B testing on onboarding and raised activation by 12 percent.",
            "Worked with engineering, design and data teams in two-week Scrum sprints.",
        ]),
        ("Product Manager", "Contoso Retail", "03/2018 – 12/2020", [
            "Launched a self-service returns flow used in 40 stores.",
            "Wrote requirements and tracked delivery in Jira.",
        ]),
        ("Business Analyst", "Fabrikam Bank", "2016 – 2018", [
            "Analysed payment data with SQL and built weekly dashboards.",
        ]),
    ],
    "education": [("BSc Computer Engineering", "Example State University", "2012 – 2016")],
    "skills": ["Jira", "SQL", "Python", "Figma", "Scrum", "A/B testing", "Roadmapping"],
    "languages": ["English (native)", "Spanish (B2)"],
}

TR = {
    "name": "Deniz Yılmaz",
    "title": "Yazılım Geliştirici",
    "email": "deniz.yilmaz@example.com",
    "phone": "+90 555 010 20 30",
    "linkedin": "linkedin.com/in/deniz-yilmaz-example",
    "headings": {"summary": "Özet", "experience": "İş Deneyimi", "education": "Eğitim", "skills": "Yetenekler", "languages": "Diller"},
    "summary": "İstanbul'da web ve mobil projelerde altı yıllık deneyime sahip yazılım geliştirici.",
    "jobs": [
        ("Kıdemli Yazılım Geliştirici", "Örnek Teknoloji A.Ş.", "Oca 2021 – Günümüz", [
            "Python ve Django ile ödeme altyapısını geliştirdim.",
            "Kubernetes üzerinde çalışan mikroservislerin bakımını üstlendim.",
            "TÜBİTAK destekli bir projede görüntü işleme modülü yazdım.",
        ]),
        ("Yazılım Geliştirici", "Deneme Yazılım Ltd.", "06/2019 – 12/2020", [
            "React ile müşteri paneli geliştirdim.",
            "PostgreSQL sorgularını optimize ettim.",
        ]),
    ],
    "education": [("Bilgisayar Mühendisliği, Lisans", "Örnek Üniversitesi", "2015 – 2019")],
    "skills": ["Python", "Django", "React", "PostgreSQL", "Docker", "Kubernetes", "Git"],
    "languages": ["Türkçe (ana dil)", "İngilizce (C1)"],
}


# HTML resumes, printed by Chrome ------------------------------------------------------------------

BASE_CSS = """
@page { size: A4; margin: 16mm; }
body { font-family: Helvetica, Arial, sans-serif; font-size: 10.5pt; line-height: 1.4; color: #111; margin: 0; }
h1 { font-size: 22pt; margin: 0; } .role { font-size: 12pt; margin: 2pt 0 6pt; }
h2 { font-size: 12.5pt; margin: 14pt 0 4pt; border-bottom: 1px solid #999; }
.job { margin-bottom: 6pt; } .job-head { display: flex; justify-content: space-between; font-weight: bold; }
ul { margin: 2pt 0 0 14pt; padding: 0; } p { margin: 2pt 0; }
"""


def contact_html(p, email=True, linkedin_text=True):
    parts = []
    if email:
        parts.append(p["email"])
    parts.append(p["phone"])
    parts.append(p["linkedin"] if linkedin_text else f'<a href="https://www.{p["linkedin"]}">LinkedIn</a>')
    return " · ".join(parts)


def section_html(p, key, heading=None, date_override=None):
    h = heading or p["headings"][key]
    if key == "summary":
        return f"<h2>{h}</h2><p>{p['summary']}</p>"
    if key == "experience":
        jobs = []
        for i, (title, company, dates, bullets) in enumerate(p["jobs"]):
            dates = (date_override or {}).get(i, dates)
            items = "".join(f"<li>{b}</li>" for b in bullets)
            jobs.append(f'<div class="job"><div class="job-head"><span>{title}, {company}</span><span>{dates}</span></div><ul>{items}</ul></div>')
        return f"<h2>{h}</h2>" + "".join(jobs)
    if key == "education":
        return f"<h2>{h}</h2>" + "".join(f'<div class="job-head"><span>{d}, {s}</span><span>{y}</span></div>' for d, s, y in p["education"])
    if key == "skills":
        return f"<h2>{h}</h2><p>{', '.join(p['skills'])}</p>"
    if key == "languages":
        return f"<h2>{h}</h2><p>{', '.join(p['languages'])}</p>"
    raise KeyError(key)


def page_html(body, css="", lang="en"):
    return f'<!doctype html><html lang="{lang}"><head><meta charset="utf-8"><style>{BASE_CSS}{css}</style></head><body>{body}</body></html>'


def header_html(p, **kw):
    return f"<h1>{p['name']}</h1><div class='role'>{p['title']}</div><div>{contact_html(p, **kw)}</div>"


def single_column(p, headings=None, date_override=None, **kw):
    keys = ["summary", "experience", "education", "skills", "languages"]
    return header_html(p, **kw) + "".join(section_html(p, k, (headings or {}).get(k), date_override) for k in keys)


INTER_CSS = f"""
@font-face {{ font-family: InterLocal; src: url("file://{FONTS}/inter-latin-wght-normal.woff2"); unicode-range: U+0000-00FF; }}
@font-face {{ font-family: InterLocal; src: url("file://{FONTS}/inter-latin-ext-wght-normal.woff2"); unicode-range: U+0100-024F; }}
body {{ font-family: InterLocal, sans-serif; }}
"""


def html_fixtures():
    sidebar = f"""
    <div class="grid"><aside>{header_html(EN)}{section_html(EN, 'skills')}{section_html(EN, 'languages')}</aside>
    <main>{section_html(EN, 'summary')}{section_html(EN, 'experience')}{section_html(EN, 'education')}</main></div>"""
    equal = f"""{header_html(EN)}<div class="grid2"><div>{section_html(EN, 'summary')}{section_html(EN, 'experience')}</div>
    <div>{section_html(EN, 'education')}{section_html(EN, 'skills')}{section_html(EN, 'languages')}</div></div>"""
    table = f"""{header_html(EN)}<table><tr><td class="l">{section_html(EN, 'experience')}</td>
    <td class="r">{section_html(EN, 'education')}{section_html(EN, 'skills')}{section_html(EN, 'languages')}</td></tr></table>"""
    tiny = single_column(EN).replace("<li>", '<li><span class="ghost">•</span>')
    longer = dict(EN, jobs=EN["jobs"] * 3)
    return {
        "html_single_column": (page_html(single_column(EN)), "en"),
        "html_two_pages_inter": (page_html(single_column(longer), INTER_CSS), "en"),
        "html_sidebar_left": (page_html(sidebar, ".grid{display:grid;grid-template-columns:32% 1fr;gap:18pt} aside{font-size:9.5pt}"), "en"),
        "html_two_equal_columns": (page_html(equal, ".grid2{display:grid;grid-template-columns:1fr 1fr;gap:20pt}"), "en"),
        "html_table_layout": (page_html(table, "table{width:100%;border-collapse:collapse} td{vertical-align:top} .l{width:62%;padding-right:18pt}"), "en"),
        "html_letterspaced_headings": (page_html(single_column(EN), "h2{text-transform:uppercase;letter-spacing:.45em;font-size:10pt}"), "en"),
        # As on a real resume: the visible bullet is a shape, plus a tiny text bullet for copy and paste.
        "html_tiny_text_bullets": (page_html(tiny, "ul{list-style:none;margin-left:0} li{position:relative;padding-left:12pt} li::after{content:'';position:absolute;left:2pt;top:5pt;width:4pt;height:4pt;border-radius:50%;background:#111} .ghost{font-size:.75pt;display:inline-block;width:0;overflow:visible}"), "en"),
        "html_turkish": (page_html(single_column(TR), INTER_CSS, lang="tr"), "tr"),
        "html_nonstandard_headings_dates": (page_html(single_column(
            EN, headings={"experience": "My Journey", "skills": "Toolbox", "summary": "Hello There"},
            date_override={0: "Summer 2021 – Present", 1: "Q2 2018 – Q4 2020"})), "en"),
        "html_no_email_hidden_linkedin": (page_html(single_column(EN, email=False, linkedin_text=False)), "en"),
    }


def print_pdf(html: str, target: Path, workdir: Path) -> None:
    source = workdir / (target.stem + ".html")
    source.write_text(html, encoding="utf-8")
    subprocess.run(
        [CHROME, "--headless=new", "--disable-gpu", "--no-pdf-header-footer", "--allow-file-access-from-files",
         f"--user-data-dir={workdir / 'chrome'}", f"--print-to-pdf={target}", source.as_uri()],
        check=True, capture_output=True,
    )


def scan(pdf: Path, target: Path) -> None:
    """A page image saved as PDF, like a scanned or photographed resume."""
    document = pypdfium2.PdfDocument(str(pdf))
    images = [page.render(scale=150 / 72).to_pil().convert("L") for page in document]
    images[0].save(target, "PDF", resolution=150, save_all=True, append_images=images[1:])


# Word resumes ------------------------------------------------------------------------------------

TEXT_BOX = """<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
 xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"
 xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
 xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
 xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">
<mc:AlternateContent><mc:Choice Requires="wps"><w:drawing>
<wp:anchor distT="0" distB="0" distL="114300" distR="114300" simplePos="0" relativeHeight="1" behindDoc="0" locked="0" layoutInCell="1" allowOverlap="1">
<wp:simplePos x="0" y="0"/><wp:positionH relativeFrom="column"><wp:posOffset>3600000</wp:posOffset></wp:positionH>
<wp:positionV relativeFrom="paragraph"><wp:posOffset>0</wp:posOffset></wp:positionV><wp:extent cx="2200000" cy="900000"/>
<wp:effectExtent l="0" t="0" r="0" b="0"/><wp:wrapSquare wrapText="bothSides"/><wp:docPr id="1" name="Box"/><wp:cNvGraphicFramePr/>
<a:graphic><a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"><wps:wsp><wps:cNvSpPr txBox="1"/>
<wps:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="2200000" cy="900000"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></wps:spPr>
<wps:txbx><w:txbxContent>{paragraphs}</w:txbxContent></wps:txbx><wps:bodyPr/></wps:wsp></a:graphicData></a:graphic></wp:anchor>
</w:drawing></mc:Choice></mc:AlternateContent></w:r>"""


def word_document(p, contact_in_header=False, text_box=False, columns=1, table_layout=False):
    document = docx.Document()
    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    contact = f"{p['email']} | {p['phone']} | {p['linkedin']}"
    name = document.add_paragraph()
    run = name.add_run(p["name"])
    run.bold, run.font.size = True, Pt(20)
    document.add_paragraph(p["title"])
    if contact_in_header:
        document.sections[0].header.paragraphs[0].text = contact
    else:
        document.add_paragraph(contact)

    def experience(container):
        container.add_paragraph(p["headings"]["experience"], style="Heading 1")
        for title, company, dates, bullets in p["jobs"]:
            head = container.add_paragraph()
            head.add_run(f"{title}, {company}").bold = True
            head.add_run(f"    {dates}")
            for b in bullets:
                container.add_paragraph(b, style="List Bullet")

    def rest(container, skip_skills=False):
        container.add_paragraph(p["headings"]["education"], style="Heading 1")
        for degree, school, years in p["education"]:
            container.add_paragraph(f"{degree}, {school}    {years}")
        if not skip_skills:
            container.add_paragraph(p["headings"]["skills"], style="Heading 1")
            container.add_paragraph(", ".join(p["skills"]))
        container.add_paragraph(p["headings"]["languages"], style="Heading 1")
        container.add_paragraph(", ".join(p["languages"]))

    document.add_heading(p["headings"]["summary"], level=1)
    summary = document.add_paragraph(p["summary"])
    if text_box:
        lines = [p["headings"]["skills"], ", ".join(p["skills"])]
        paragraphs = "".join(f"<w:p><w:r><w:t xml:space='preserve'>{line}</w:t></w:r></w:p>" for line in lines)
        summary._p.append(parse_xml(TEXT_BOX.format(paragraphs=paragraphs)))
    if table_layout:
        table = document.add_table(rows=1, cols=2)
        left, right = table.cell(0, 0), table.cell(0, 1)
        left.paragraphs[0].text = ""
        experience(left)
        rest(right)
    else:
        if columns > 1:
            section = document.add_section(0)  # continuous section break
            section._sectPr.find(qn("w:cols")).set(qn("w:num"), str(columns))
        experience(document)
        rest(document, skip_skills=text_box)
    return document


def word_fixtures():
    return {
        "word_single_column": word_document(EN),
        "word_contact_in_header": word_document(EN, contact_in_header=True),
        "word_text_box_skills": word_document(EN, text_box=True),
        "word_two_columns": word_document(EN, columns=2),
        "word_table_layout": word_document(EN, table_layout=True),
        "word_turkish": word_document(TR),
    }


def export_with_word(docx_path: Path, pdf_path: Path) -> None:
    WORD_DIR.mkdir(parents=True, exist_ok=True)
    inside_docx, inside_pdf = WORD_DIR / docx_path.name, WORD_DIR / pdf_path.name
    shutil.copy(docx_path, inside_docx)
    script = f'''
    tell application "Microsoft Word"
      open file name (POSIX file "{inside_docx}" as text)
      delay 1
      save as active document file name (POSIX file "{inside_pdf}" as text) file format format PDF
      close active document saving no
    end tell'''
    subprocess.run(["osascript", "-e", script], check=True, capture_output=True)
    shutil.move(inside_pdf, pdf_path)
    inside_docx.unlink()


def main() -> int:
    if not Path(CHROME).exists():
        print("Google Chrome is required.", file=sys.stderr)
        return 1
    OUT.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as tmp:
        work = Path(tmp)
        for name, (html, _) in html_fixtures().items():
            print_pdf(html, OUT / f"{name}.pdf", work)
            print("made", name)
        scan(OUT / "html_single_column.pdf", OUT / "scanned_single_column.pdf")
        print("made scanned_single_column")
        for name, document in word_fixtures().items():
            document.save(OUT / f"{name}.docx")
            export_with_word(OUT / f"{name}.docx", OUT / f"{name}.pdf")
            print("made", name, "(docx and pdf)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
