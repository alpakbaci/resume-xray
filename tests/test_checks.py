from pathlib import Path

import docx
import pytest
from docx.oxml import parse_xml
from pdfgen import Text, write_pdf
from test_docx_reader import TEXT_BOX

from resume_xray.checks import check_garbled, run_checks
from resume_xray.config import load_weights
from resume_xray.models import Document, Page
from resume_xray.readers import read_document
from resume_xray.text import fold

SYNTHETIC = Path(__file__).parent / "fixtures" / "synthetic"
CONTACT = Text(50, 80, "jane.doe@example.com  +90 555 123 45 67")


def clean_resume(extra=()):
    return [
        Text(50, 50, "Jane Doe", size=20, bold=True),
        CONTACT,
        Text(50, 120, "EXPERIENCE", size=13, bold=True),
        Text(50, 140, "Backend Developer, Acme Corp, 01/2021 - Present"),
        Text(50, 180, "EDUCATION", size=13, bold=True),
        Text(50, 200, "BSc Computer Engineering, 2016 - 2020"),
        Text(50, 240, "SKILLS", size=13, bold=True),
        Text(50, 260, "Python, Docker, PostgreSQL"),
        *extra,
    ]


def report_for(tmp_path, pages):
    return run_checks(read_document(write_pdf(tmp_path / "cv.pdf", pages)))


def checks(report):
    return {(f.check, f.title) for f in report.findings}


def test_clean_resume_is_ok(tmp_path):
    report = report_for(tmp_path, [clean_resume()])

    assert report.findings == []
    assert report.verdict == "OK"
    assert set(report.sections) == {"experience", "education", "skills"}
    assert report.fields == {"email": "jane.doe@example.com", "phone": "+90 555 123 45 67"}


def test_scanned_page_is_unreadable():
    report = run_checks(read_document(SYNTHETIC / "scanned_single_column.pdf"))

    finding = next(f for f in report.findings if f.check == "2.1")
    assert finding.severity == "critical" and finding.pages == (1,)
    assert report.verdict == "Unreadable"


def test_short_last_page_is_not_a_scan(tmp_path):
    # Found on a generated two-page resume: page 2 held only the Languages section.
    report = report_for(tmp_path, [clean_resume(), [Text(50, 60, "LANGUAGES", size=13, bold=True), Text(50, 80, "English")]])

    assert not any(f.check == "2.1" for f in report.findings)


def test_two_columns_scramble_position_order(tmp_path):
    left = [Text(50, 120 + 16 * i, f"Left column line {i} about backend work" + " and APIs" * (i % 2)) for i in range(12)]
    right = [Text(330, 120 + 16 * i, f"Right item {i} skills") for i in range(12)]
    report = report_for(tmp_path, [[Text(50, 50, "Jane Doe", size=20, bold=True), CONTACT, *left, *right]])

    finding = next(f for f in report.findings if f.check == "2.3")
    assert finding.severity == "warning"
    assert "position" in finding.strategies
    assert "columns" not in finding.strategies


def test_tiny_bullets_are_reported_for_position_order(tmp_path):
    bullets = [
        item
        for i in range(3)
        # As on the real resume: a 0.75 pt bullet on the text's baseline, over its first letter.
        for item in (Text(58, 300 + 17 * i, "•", size=0.75), Text(58, 300 + 17 * i, f"Did thing number {i}", size=10))
    ]
    report = report_for(tmp_path, [clean_resume(bullets)])

    finding = next(f for f in report.findings if f.check == "2.8")
    assert finding.strategies == ("position",)
    assert "tiny glyphs" in finding.detail


def test_missing_email_is_critical(tmp_path):
    pages = [[t for t in clean_resume() if t is not CONTACT]]
    report = report_for(tmp_path, pages)

    assert ("2.7", "No email address found") in checks(report)
    assert report.verdict == "Unreadable"


def test_vague_dates_and_unknown_headings(tmp_path):
    extra = [
        Text(50, 160, "Intern, Beta Ltd, Summer 2019"),
        Text(50, 300, "MY JOURNEY", size=13, bold=True),
        Text(50, 320, "Travelled a lot"),
    ]
    report = report_for(tmp_path, [clean_resume(extra)])

    dates = next(f for f in report.findings if f.title == "Dates that cannot be parsed")
    assert dates.examples == ("Summer 2019",)
    heading = next(f for f in report.findings if f.title == "Section heading not recognized")
    assert heading.examples == ("MY JOURNEY",)


def test_missing_core_section(tmp_path):
    pages = [[t for t in clean_resume() if t.text not in ("SKILLS", "Python, Docker, PostgreSQL")]]
    report = report_for(tmp_path, pages)

    finding = next(f for f in report.findings if f.title == "Standard section not found")
    assert finding.examples == ("Skills",)


@pytest.fixture
def docx_resume(tmp_path):
    def build(contact_in_header: bool, text_box: bool):
        document = docx.Document()
        if contact_in_header:
            document.sections[0].header.paragraphs[0].text = "jane.doe@example.com"
        else:
            document.add_paragraph("jane.doe@example.com")
        for heading, body in (("Experience", "Developer, 2020 - 2023"), ("Education", "BSc"), ("Skills", "Python")):
            document.add_heading(heading, level=1)
            paragraph = document.add_paragraph(body)
        if text_box:
            paragraph._p.append(parse_xml(TEXT_BOX))
        path = tmp_path / "cv.docx"
        document.save(path)
        return run_checks(read_document(path))

    return build


def test_docx_contact_in_header_and_text_box(docx_resume):
    report = docx_resume(contact_in_header=True, text_box=True)

    titles = {f.title: f for f in report.findings if f.check == "2.5"}
    assert titles["Contact details only in the header or footer"].strategies == ("body",)
    assert titles["Text inside text boxes"].examples == ("Languages: English, Turkish",)
    # The full reading still finds the email, so it is not reported as missing.
    assert ("2.7", "No email address found") not in checks(report)


def test_docx_clean(docx_resume):
    report = docx_resume(contact_in_header=False, text_box=False)

    assert report.findings == []


def test_garbled_characters():
    text = "Jane (cid:12)(cid:13)(cid:14) Doe  Python developer"
    document = Document("cv.pdf", "pdf", [Page(1, 595, 842, readings={"columns": text})])

    [finding] = check_garbled(document, load_weights()["parseability"])
    assert finding.severity == "critical"


def test_fold_matches_across_languages_and_case():
    assert fold("DENEYİM") == fold("Deneyim") == "deneyim"
    assert fold("EXPERIENCE") == fold("Experience")
    assert fold("Eğitim:") == "egitim"


def test_narrow_second_column_glued_to_lines_is_caught(tmp_path):
    # One short skill per row: word order barely changes, but position order glues each skill to
    # the end of an experience line.
    left = [Text(50, 130 + 16 * i, f"Experience line number {i} about delivering products" + " fast" * (i % 3)) for i in range(8)]
    right = [Text(380, 130 + 16 * i, skill) for i, skill in enumerate(["Jira", "Figma", "SQL", "Agile", "Scrum", "Excel", "Git", "Python"])]
    report = report_for(tmp_path, [[Text(50, 50, "Jane Doe", size=20, bold=True), CONTACT, *left, *right]])

    finding = next(f for f in report.findings if f.check == "2.3")
    assert finding.code == "columns_scramble" and "position" in finding.strategies
