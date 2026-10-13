import docx
import pytest
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx.shared import Pt

from resume_xray.cli import format_raw
from resume_xray.readers import UnsupportedFormat, read_document

TEXT_BOX = """
<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
     xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"
     xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"
     xmlns:v="urn:schemas-microsoft-com:vml">
  <mc:AlternateContent>
    <mc:Choice Requires="wps"><w:drawing><wps:txbx><w:txbxContent>
      <w:p><w:r><w:t>Languages: English, Turkish</w:t></w:r></w:p>
    </w:txbxContent></wps:txbx></w:drawing></mc:Choice>
    <mc:Fallback><w:pict><v:textbox><w:txbxContent>
      <w:p><w:r><w:t>Languages: English, Turkish</w:t></w:r></w:p>
    </w:txbxContent></v:textbox></w:pict></mc:Fallback>
  </mc:AlternateContent>
</w:r>
"""


@pytest.fixture
def resume_docx(tmp_path):
    document = docx.Document()
    document.sections[0].header.paragraphs[0].text = "jane.doe@example.com | +90 555 123 45 67"
    name = document.add_paragraph().add_run("Jane Doe")
    name.bold = True
    name.font.size = Pt(20)
    document.add_heading("Experience", level=1)
    document.add_paragraph("Backend Developer, Acme Corp, 01/2021 - Present")
    box_holder = document.add_paragraph("Summary")
    box_holder._p.append(parse_xml(TEXT_BOX))
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Python"
    table.cell(0, 1).text = "Docker"
    # The default template already declares <w:cols>; make it two columns.
    document.sections[0]._sectPr.find(qn("w:cols")).set(qn("w:num"), "2")
    path = tmp_path / "cv.docx"
    document.save(path)
    return path


def test_reads_regions_fonts_and_columns(resume_docx):
    document = read_document(resume_docx)

    assert document.kind == "docx"
    page = document.pages[0]
    by_region = {}
    for line in page.lines:
        by_region.setdefault(line.region, []).append(line.text)

    assert by_region["header"] == ["jane.doe@example.com | +90 555 123 45 67"]
    assert by_region["table"] == ["Python", "Docker"]
    # Word stores a text box twice; it must be read once.
    assert by_region["textbox"] == ["Languages: English, Turkish"]
    assert "Languages" not in page.text
    assert "jane.doe" not in page.text

    name = next(line for line in page.lines if line.text == "Jane Doe")
    assert name.bold and name.font_size == 20
    heading = next(line for line in page.lines if line.text == "Experience")
    assert heading.style == "Heading 1" and heading.bold
    assert page.columns == 2


def test_raw_output_labels_regions(resume_docx):
    raw = format_raw(read_document(resume_docx))

    assert "── Body ──" in raw
    assert "── Header (some ATSs skip this) ──" in raw
    assert "── Text boxes (some ATSs skip or misplace these) ──" in raw


@pytest.mark.parametrize(("name", "message"), [("cv.doc", "Old .doc"), ("cv.pages", "Unsupported")])
def test_unsupported_formats(tmp_path, name, message):
    with pytest.raises(UnsupportedFormat, match=message):
        read_document(tmp_path / name)
