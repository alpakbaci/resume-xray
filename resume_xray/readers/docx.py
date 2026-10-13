"""DOCX reader: python-docx plus raw XML for text boxes (ALGORITHM.md §2.0)."""

from pathlib import Path

import docx
from docx.document import Document as DocxDocument
from docx.table import Table
from docx.text.paragraph import Paragraph

from resume_xray.models import Document, Line, Page, Region

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
EMU_PER_PT = 12700


def read_docx(path: Path) -> Document:
    document = docx.Document(str(path))
    lines: list[Line] = []

    for block in _iter_body(document):
        if isinstance(block, Paragraph):
            lines.extend(_paragraph_lines(block, "body"))
            lines.extend(_textbox_lines(block._p))
        else:
            lines.extend(_table_lines(block))

    seen_parts = set()
    for section in document.sections:
        for region, parts in (
            ("header", (section.header, section.first_page_header, section.even_page_header)),
            ("footer", (section.footer, section.first_page_footer, section.even_page_footer)),
        ):
            for part in parts:
                if part.is_linked_to_previous or id(part.part) in seen_parts:
                    continue
                seen_parts.add(id(part.part))
                for paragraph in part.paragraphs:
                    lines.extend(_paragraph_lines(paragraph, region))
                for table in part.tables:
                    lines.extend(_table_lines(table, region))

    section = document.sections[0] if document.sections else None
    page = Page(
        number=1,
        width=section.page_width / EMU_PER_PT if section and section.page_width else None,
        height=section.page_height / EMU_PER_PT if section and section.page_height else None,
        text="\n".join(line.text for line in lines if line.region in ("body", "table")),
        lines=lines,
        char_count=sum(1 for line in lines for c in line.text if not c.isspace()),
        columns=max((_declared_columns(s) for s in document.sections), default=1),
    )
    return Document(name=path.name, kind="docx", pages=[page])


def _iter_body(document: DocxDocument):
    """Paragraphs and tables in the order they appear in the body."""
    for child in document.element.body.iterchildren():
        if child.tag == f"{{{W}}}p":
            yield Paragraph(child, document)
        elif child.tag == f"{{{W}}}tbl":
            yield Table(child, document)


def _paragraph_lines(paragraph: Paragraph, region: Region) -> list[Line]:
    text = paragraph.text.strip()
    if not text:
        return []
    style = paragraph.style
    runs = [r for r in paragraph.runs if r.text.strip()]
    sizes = [r.font.size.pt for r in runs if r.font.size]
    font_size = max(sizes) if sizes else _style_size(style)
    bold = bool(runs) and all(_run_bold(r, style) for r in runs)
    return [
        Line(text=line, page=1, region=region, font_size=font_size, bold=bold, style=style.name if style else None)
        for line in (part.strip() for part in text.split("\n"))
        if line
    ]


def _table_lines(table: Table, region: Region = "table") -> list[Line]:
    lines = []
    for row in table.rows:
        seen_cells = set()
        for cell in row.cells:
            # Merged cells are returned once per grid column; read each only once.
            if id(cell._tc) in seen_cells:
                continue
            seen_cells.add(id(cell._tc))
            for paragraph in cell.paragraphs:
                lines.extend(_paragraph_lines(paragraph, region))
    return lines


def _textbox_lines(element) -> list[Line]:
    """Text boxes inside a paragraph. Word stores each box twice (DrawingML and a VML fallback);
    the copy under mc:Fallback is skipped."""
    lines = []
    for box in element.iter(f"{{{W}}}txbxContent"):
        if any(a.tag == f"{{{MC}}}Fallback" for a in box.iterancestors()):
            continue
        for p in box.iter(f"{{{W}}}p"):
            text = "".join(t.text or "" for t in p.iter(f"{{{W}}}t")).strip()
            if text:
                lines.append(Line(text=text, page=1, region="textbox"))
    return lines


def _run_bold(run, style) -> bool:
    if run.bold is not None:
        return run.bold
    while style is not None:
        if style.font.bold is not None:
            return style.font.bold
        style = style.base_style
    return False


def _style_size(style) -> float | None:
    while style is not None:
        if style.font.size:
            return style.font.size.pt
        style = style.base_style
    return None


def _declared_columns(section) -> int:
    cols = section._sectPr.find(f"{{{W}}}cols")
    if cols is None:
        return 1
    return int(cols.get(f"{{{W}}}num", "1"))
